"""The operations dataset — every operation, one row per truck load, in the
layout management sent as "DATA SET.xlsx" (less SURVEYOR, which management
said to drop).

Each row carries the full chain for one load: licence (BFL, Naval
Clearance), sourcing (vendor, PFI), the truck and its waybill, its dates,
safety checks and issues, and the litres it loaded and discharged. The
vessel-delivery columns (receiving vessel through BDR QTY) describe the
operation as a whole, so they sit on the operation's FIRST row only — repeat
them on every truck row and a column total in Excel counts each delivery
once per truck.

Approved ship-to-ship collections get a row each too (LOADING PORT reads
"STS from <vessel>"). An operation with neither trucks nor collections still
gets one row, so it is not missing from the dataset.

Units follow the system rule: truck quantities are litres (the truck
`*_mt` columns hold litres despite their name), vessel quantities MT(vac).
Dates are Lagos local time. Unknown values are left blank, never guessed.

Nothing here writes.
"""
from __future__ import annotations

import io
from collections import defaultdict
from typing import Any, Dict, List, Optional

from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

COLUMNS: List[tuple] = [
    ("A", "BFLA"),
    ("B", "NAVAL CLEARANCE"),
    ("C", "PRODUCT"),
    ("D", "VENDORS"),
    ("E", "PFI"),
    ("F", "TRUCK NO"),
    ("G", "WAY BILL NO."),
    ("H", "WAY BILL TRUCK NO."),
    ("I", "LOADING PORT"),
    ("J", "DRIVER'S NAME"),
    ("K", "DRIVER'S NUMBER"),
    ("L", "DATE LOADED"),
    ("M", "DEPARTURE DATE"),
    ("N", "SAFETY CHECK"),
    ("O", "ISSUES"),
    ("P", "DATE ARRIVED AT DISCHARGE PORT"),
    ("Q", "DISCHARGED DATE"),
    ("R", "BUNKER VESSEL"),
    ("S", "LOADED QTY"),
    ("T", "DISCHARGE QTY"),
    ("U", "VARIANCE"),
    ("V", "OPERATION ID"),
    ("W", "RECEIVING VESSEL"),
    ("X", "DELIVERY LOCATION"),
    ("Y", "DELIVERY DATE"),
    ("Z", "QUANTITY (GOV)"),
    ("AA", "VCF"),
    ("AB", "DENSITY"),
    ("AC", "QUANTITY ON BOARD (MT)"),
    ("AD", "DISCHARGED QTY MT"),
    ("AE", "RECEIVED QTY MT"),
    ("AF", "VARIANCE"),
    ("AG", "BDR QTY"),
]

# Columns that belong to the operation's delivery, not to one truck.
_DELIVERY_COLUMNS = ("W", "X", "Y", "Z", "AA", "AB", "AC", "AD", "AE", "AF", "AG")

_LAGOS = "Africa/Lagos"


def _date(d) -> str:
    return d.strftime("%d/%m/%Y") if d else ""


def _num(v, dp: int = 2) -> Optional[float]:
    if v is None:
        return None
    try:
        return round(float(v), dp)
    except (TypeError, ValueError):
        return None


def _diff(a, b, dp: int) -> Optional[float]:
    if a is None or b is None:
        return None
    return round(float(a) - float(b), dp)


def _hse_summary(hse) -> str:
    """'Pre: Satisfactory; Transfer: ...' from a collection's HSE record."""
    labels = (("pre", "Pre"), ("during", "Transfer"), ("post", "Post"))
    parts = []
    for key, label in labels:
        phase = (hse or {}).get(key)
        if phase:
            ok = phase.get("result") == "satisfactory"
            parts.append(f"{label}: {'Satisfactory' if ok else 'Not satisfactory'}")
    return "; ".join(parts)


def _joined(values) -> str:
    """Distinct non-blank values, comma-joined — '0.9800' and '0.98' are
    one value, so numbers are normalised before de-duplicating."""
    seen: List[str] = []
    for v in values:
        if v is None:
            continue
        s = f"{float(v):g}" if not isinstance(v, str) else v.strip()
        if s and s not in seen:
            seen.append(s)
    return ", ".join(seen)


async def build_rows(db: AsyncSession) -> List[Dict[str, Any]]:
    """Every live operation, oldest first. Five queries in total, whatever the
    number of operations — the API sits a long round trip from the database."""
    ops = (await db.execute(text("""
        SELECT o.id, o.operation_number, o.loading_location, o.discharge_location,
               v.vessel_name AS our_vessel,
               (SELECT string_agg(DISTINCT p.product_type, ', ')
                  FROM operation_products p WHERE p.operation_id = o.id) AS products,
               (SELECT string_agg(DISTINCT n.clearance_number, ', ')
                  FROM operation_naval_clearances l
                  JOIN naval_clearances n ON n.id = l.naval_clearance_id
                 WHERE l.operation_id = o.id) AS clearances,
               (SELECT string_agg(DISTINCT b.bfl_number, ', ')
                  FROM operation_naval_clearances l
                  JOIN naval_clearance_drawdowns d ON d.naval_clearance_id = l.naval_clearance_id
                  JOIN bfls b ON b.id = d.bfl_id
                 WHERE l.operation_id = o.id) AS bfls,
               (SELECT string_agg(DISTINCT x.pfi_number, ', ') FROM (
                    SELECT f.pfi_number FROM pfi_allocations a JOIN pfis f ON f.id = a.pfi_id
                     WHERE a.operation_id = o.id
                    UNION SELECT f.pfi_number FROM pfis f WHERE f.operation_id = o.id
                ) x) AS pfis
        FROM operations o
        LEFT JOIN vessels v ON v.id = o.vessel_id
        WHERE o.deleted_at IS NULL AND o.status <> 'cancelled'
        ORDER BY o.created_at
    """))).mappings().all()

    trucks = (await db.execute(text(f"""
        SELECT t.operation_id, tr.truck_number, w.waybill_truck_number,
               t.product_type, t.vendor_name, t.driver_name, t.driver_phone,
               COALESCE(t.waybill_document_number, t.waybill_number) AS waybill,
               t.loading_location,
               COALESCE(t.destination_vessel_name, dv.vessel_name) AS bunker_vessel,
               COALESCE(t.loading_gov, t.loading_received_quantity_litres, t.quantity_loaded_mt) AS loaded,
               t.quantity_discharged_mt AS discharged,
               (COALESCE(t.loading_quantity_recorded_at, t.arrived_loading_at, t.departed_loading_at)
                    AT TIME ZONE '{_LAGOS}') AS loaded_at,
               (t.departed_loading_at AT TIME ZONE '{_LAGOS}') AS departed_at,
               (COALESCE(t.arrived_discharge_at, t.transit_end_at) AT TIME ZONE '{_LAGOS}') AS arrived_at,
               (t.discharge_end_at AT TIME ZONE '{_LAGOS}') AS discharged_at,
               (SELECT string_agg(
                        initcap(a.phase::text) || ': ' ||
                        CASE a.result::text WHEN 'satisfactory' THEN 'Satisfactory' ELSE 'Not satisfactory' END,
                        '; ' ORDER BY a.phase)
                  FROM truck_safety_audits a WHERE a.truck_op_id = t.id) AS safety,
               (SELECT string_agg(i.title, '; ' ORDER BY i.created_at)
                  FROM truck_issues i
                 WHERE i.truck_id = t.truck_id AND i.operation_id = t.operation_id) AS issues
        FROM truck_operations t
        JOIN trucks tr ON tr.id = t.truck_id
        LEFT JOIN truck_waivers w ON w.id = t.waiver_id
        LEFT JOIN vessels dv ON dv.id = t.destination_vessel_id
        WHERE t.status <> 'cancelled'
        ORDER BY COALESCE(t.departed_loading_at, t.created_at)
    """))).mappings().all()

    # Vessel BDNs only (vessel_activity_id set) — the delivery to the client.
    # Rejected ones were never accepted as the record.
    deliveries = (await db.execute(text(f"""
        SELECT b.operation_id, b.receiving_vessel, b.discharge_location,
               (COALESCE(b.discharge_completed_at::timestamptz, b.discharge_completion_date::timestamptz)
                    AT TIME ZONE '{_LAGOS}') AS delivered_at,
               COALESCE(b.received_gov, b.discharge_gov) AS gov,
               b.vcf, b.density, b.discharge_mt_vacuum, b.received_mt_vacuum, b.bdr_quantity_mt
        FROM bdns b
        WHERE b.vessel_activity_id IS NOT NULL AND b.status <> 'rejected'
        ORDER BY b.created_at
    """))).mappings().all()

    # Approved ship-to-ship collections — product the barge took on from
    # another vessel. One row each, beside the truck rows; GOV goes in
    # LOADED QTY because GOV is litres, the same unit as the truck figures.
    collections = (await db.execute(text(f"""
        SELECT c.operation_id, c.source_vessel_name, c.product_type, c.gov, c.hse,
               v.vessel_name AS barge,
               (c.cast_off_at AT TIME ZONE '{_LAGOS}') AS cast_off_at,
               (c.receiving_completed_at AT TIME ZONE '{_LAGOS}') AS completed_at
        FROM sts_collections c
        JOIN vessels v ON v.id = c.vessel_id
        WHERE c.status = 'approved'
        ORDER BY COALESCE(c.receiving_completed_at, c.created_at)
    """))).mappings().all()

    # Fuel on board just before the operation's first discharge.
    on_board = (await db.execute(text("""
        SELECT DISTINCT ON (r.operation_id) r.operation_id, r.rob_before_mt
        FROM rob_entries r
        WHERE r.operation_id IS NOT NULL AND r.entry_type = 'discharge'
        ORDER BY r.operation_id, r.created_at
    """))).mappings().all()

    trucks_by_op: Dict[Any, list] = defaultdict(list)
    for t in trucks:
        trucks_by_op[t["operation_id"]].append(t)
    deliveries_by_op: Dict[Any, list] = defaultdict(list)
    for d in deliveries:
        deliveries_by_op[d["operation_id"]].append(d)
    on_board_by_op = {r["operation_id"]: r["rob_before_mt"] for r in on_board}
    collections_by_op: Dict[Any, list] = defaultdict(list)
    for c in collections:
        collections_by_op[c["operation_id"]].append(c)

    rows: List[Dict[str, Any]] = []
    for op in ops:
        op_rows: List[Dict[str, Any]] = []
        op_trucks = trucks_by_op.get(op["id"], [])
        op_collections = collections_by_op.get(op["id"], [])
        # An operation with neither still gets one row, so it isn't missing.
        for t in op_trucks or ([] if op_collections else [None]):
            row: Dict[str, Any] = {
                "A": op["bfls"] or "",
                "B": op["clearances"] or "",
                "C": (t and t["product_type"]) or op["products"] or "",
                "I": (t and t["loading_location"]) or op["loading_location"] or "",
                "R": (t and t["bunker_vessel"]) or op["our_vessel"] or "",
                "E": op["pfis"] or "",
                "V": op["operation_number"],
            }
            if t:
                row.update({
                    "D": t["vendor_name"] or "",
                    "F": t["truck_number"] or "",
                    "G": t["waybill"] or "",
                    "H": t["waybill_truck_number"] or "",
                    "J": t["driver_name"] or "",
                    "K": t["driver_phone"] or "",
                    "L": _date(t["loaded_at"]),
                    "M": _date(t["departed_at"]),
                    "N": t["safety"] or "",
                    "O": t["issues"] or "",
                    "P": _date(t["arrived_at"]),
                    "Q": _date(t["discharged_at"]),
                    "S": _num(t["loaded"]),
                    "T": _num(t["discharged"]),
                    "U": _diff(t["loaded"], t["discharged"], 2),
                })
            op_rows.append(row)

        for c in op_collections:
            op_rows.append({
                "A": op["bfls"] or "",
                "B": op["clearances"] or "",
                "C": c["product_type"] or op["products"] or "",
                "E": op["pfis"] or "",
                "I": f"STS from {c['source_vessel_name']}",
                "L": _date(c["completed_at"]),
                "M": _date(c["cast_off_at"]),
                "N": _hse_summary(c["hse"]),
                "R": c["barge"] or op["our_vessel"] or "",
                "S": _num(c["gov"]),
                "V": op["operation_number"],
            })

        ds = deliveries_by_op.get(op["id"], [])
        discharged = [d["discharge_mt_vacuum"] for d in ds if d["discharge_mt_vacuum"] is not None]
        received = [d["received_mt_vacuum"] for d in ds if d["received_mt_vacuum"] is not None]
        govs = [d["gov"] for d in ds if d["gov"] is not None]
        bdrs = [d["bdr_quantity_mt"] for d in ds if d["bdr_quantity_mt"] is not None]
        discharged_mt = sum(discharged) if discharged else None
        received_mt = sum(received) if received else None
        delivered = [d["delivered_at"] for d in ds if d["delivered_at"]]
        op_rows[0].update({
            "W": _joined(d["receiving_vessel"] for d in ds),
            "X": _joined(d["discharge_location"] for d in ds) or op["discharge_location"] or "",
            "Y": _date(max(delivered)) if delivered else "",
            "Z": _num(sum(govs)) if govs else None,
            "AA": _joined(d["vcf"] for d in ds),
            "AB": _joined(d["density"] for d in ds),
            "AC": _num(on_board_by_op.get(op["id"]), 3),
            "AD": _num(discharged_mt, 3),
            "AE": _num(received_mt, 3),
            "AF": _diff(discharged_mt, received_mt, 3),
            # The client's BDR figure, summed over the operation's deliveries.
            "AG": _num(sum(bdrs), 3) if bdrs else None,
        })
        rows.extend(op_rows)

    return rows


def to_xlsx(rows: List[Dict[str, Any]]) -> bytes:
    """One sheet, the headings, frozen header row."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "DATA SET"

    header_font = Font(bold=True, color="FFFFFF", size=9)
    header_fill = PatternFill("solid", fgColor="0A1F3C")
    for idx, (_, heading) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=idx, value=heading)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(idx)].width = min(max(len(heading) * 0.9, 12), 30)
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"

    for r, row in enumerate(rows, start=2):
        for idx, (letter, _) in enumerate(COLUMNS, start=1):
            value = row.get(letter)
            ws.cell(row=r, column=idx, value="" if value is None else value)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
